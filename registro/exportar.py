"""Exportar reportes: PDF, HTML, Excel y CSV con las mismas métricas que ve el usuario.

`armar()` junta los datos una sola vez (con filtros de privacidad ya aplicados) y
cada `a_*()` los escribe en su formato. Nada acá toca la base ni la interfaz.
"""

import csv
import os
import tempfile
from datetime import datetime

from . import __version__, analisis as A, metricas as M

SECCIONES = [
    ("resumen", "Resumen y métricas clave"),
    ("riesgo", "Riesgo: drawdown, rachas y costos"),
    ("curva", "Curva de equity"),
    ("dias", "Resultado por día"),
    ("hora", "Por hora de entrada"),
    ("activo", "Por activo"),
    ("semana", "Por día de la semana"),
    ("setup", "Por setup (playbook)"),
    ("ops", "Detalle de operaciones"),
]
FORMATOS = [("pdf", "PDF", ".pdf"), ("html", "HTML", ".html"), ("xlsx", "Excel", ".xlsx"), ("csv", "CSV", ".csv")]
TEMAS = [("blanco", "Blanco (para imprimir)"), ("papel", "Papel (marca gabomondo)"), ("oscuro", "Oscuro")]

MARCA_TEXTO = "Hecho con Registro de operativa · gabomondo.com"

OPCIONES = dict(secciones=[s for s, _ in SECCIONES], comentarios=False, anonimizar=False,
                marca=True, tema="blanco", fuentes=True, csv_sep=",")

# Paleta de marca (styles.css de gabomondo.com). `pos`/`neg` cambian en oscuro para mantener contraste.
PALETAS = {
    "papel": dict(bg="#F5F1EA", card="#FBF9F5", text="#0E1116", muted="#5C6671", line="#DDD6CB",
                  signal="#2C5AA0", pos="#3A6B3A", neg="#8B3A2E", head_bg="#0E1116", head_fg="#F5F1EA"),
    "blanco": dict(bg="#FFFFFF", card="#FFFFFF", text="#0E1116", muted="#5C6671", line="#DDD6CB",
                   signal="#2C5AA0", pos="#3A6B3A", neg="#8B3A2E", head_bg="#0E1116", head_fg="#F5F1EA"),
    "oscuro": dict(bg="#0E1116", card="#151A21", text="#F5F1EA", muted="#A8A398", line="#2A3039",
                   signal="#6FA3E8", pos="#72B272", neg="#D4756B", head_bg="#F5F1EA", head_fg="#0E1116"),
}


class ExportError(Exception):
    """Error apto para mostrarle al usuario."""


# ---------------------------------------------------------------- formato
def usd(x, signo=True, dec=2):
    if x is None:
        return "—"
    return f"{x:+,.{dec}f}" if signo else f"{x:,.{dec}f}"


def num(x, dec=2):
    return "—" if x is None else f"{x:,.{dec}f}"


def tono(x):
    return "pos" if x and x > 0 else ("neg" if x and x < 0 else "neu")


def _f(fecha):
    """'2026-10-15' -> '15/10/2026'."""
    try:
        return datetime.strptime(fecha, "%Y-%m-%d").strftime("%d/%m/%Y")
    except (ValueError, TypeError):
        return fecha or ""


# ---------------------------------------------------------------- datos
def _ops_normalizadas(rows, opciones):
    ops = [o for o in rows if o.get("profit") is not None and o.get("fecha")]
    ops.sort(key=lambda o: (o["fecha"], o.get("hora") or ""))
    alias, out = {}, []
    for o in ops:
        emp, cta = (o.get("empresa") or "").strip(), (o.get("cuenta") or "").strip()
        if opciones.get("anonimizar"):
            k = (emp, cta)
            if k not in alias:
                n = len(alias)
                alias[k] = "Cuenta " + (chr(65 + n) if n < 26 else str(n + 1))
            emp, cta = "", alias[k]
        out.append(dict(
            fecha=o["fecha"], hora=(o.get("hora") or ""), activo=(o.get("activo") or ""),
            lado=(o.get("lado") or ""), contratos=o.get("contratos"), setup=(o.get("tipo") or ""),
            bruto=o.get("bruto"), costo=o.get("comision"), neto=float(o["profit"]),
            empresa=emp, cuenta=cta, emo=o.get("emo"),
            comentario=((o.get("comentario") or "").strip() if opciones.get("comentarios") else "")))
    return out


def armar(rows, titulo="Reporte de operativa", periodo="", filtros="", opciones=None):
    """Junta todo lo que necesitan los escritores. `rows` son filas de db_all()."""
    op = dict(OPCIONES)
    op.update(opciones or {})
    sec = [s for s, _ in SECCIONES if s in set(op["secciones"])]
    validas = [o for o in rows if o.get("profit") is not None and o.get("fecha")]
    D = dict(titulo=titulo or "Reporte de operativa", periodo=periodo, filtros=filtros, opciones=op,
             secciones=sec, n=len(validas), generado=datetime.now().strftime("%d/%m/%Y %H:%M"),
             version=__version__, paleta=PALETAS.get(op["tema"], PALETAS["blanco"]))
    if not validas:
        return D
    D["resumen"] = A.resumen(validas)
    D["extra"] = M.extra(validas)
    D["curva"] = [v for _, _, v in M.curva(validas)]
    D["curva_fechas"] = [f for f, _, _ in M.curva(validas)]
    D["dias"] = A.por_dia(validas)                                   # del más nuevo al más viejo
    D["hora"] = A.por_hora(validas)
    D["activo"] = A.por_activo(validas)
    D["semana"] = A.por_dia_semana(validas)
    D["setup"] = A._agrupar(validas, lambda o: (o.get("tipo") or "").strip() or None,
                            orden=lambda x: -x["neto"])
    D["ops"] = _ops_normalizadas(validas, op)
    return D


def kpis_resumen(D):
    """[(etiqueta, valor, detalle, tono)] de la sección Resumen."""
    r = D["resumen"]
    return [
        ("Neto", usd(r["neto"], dec=0), f"{r['n']} trades · {r['trades_por_dia']:.1f} por día", tono(r["neto"])),
        ("Win rate", f"{r['wr']:.0f}%", f"{r['dias_verdes']} de {r['dias']} días verdes ({r['pct_dias_verdes']:.0f}%)", "neu"),
        ("Profit factor", "—" if r["profit_factor"] is None else f"{r['profit_factor']:.2f}", "ganado ÷ perdido", "neu"),
        ("Expectativa", usd(r["expectativa"], dec=1), "USD por trade", tono(r["expectativa"])),
        ("Ganador / perdedor", f"{r['prom_ganador']:+,.0f} / {r['prom_perdedor']:+,.0f}",
         "" if r["payoff"] is None else f"payoff {r['payoff']:.2f}", "neu"),
        ("Mejor / peor día", f"{r['mejor_dia']['neto']:+,.0f} / {r['peor_dia']['neto']:+,.0f}",
         f"promedio {r['prom_dia']:+,.0f} por día", "neu"),
    ]


def kpis_riesgo(D):
    e = D["extra"]
    dd_det = ""
    if e["dd_max"] > 0 and e["dd_hasta"]:
        dd_det = f"{_f(e['dd_desde']) if e['dd_desde'] else 'inicio'} → {_f(e['dd_hasta'])}"
    return [
        ("Drawdown máximo", usd(-e["dd_max"], dec=0) if e["dd_max"] else "0", dd_det or "sin caídas", "neg" if e["dd_max"] else "neu"),
        ("Drawdown actual", usd(-e["dd_actual"], dec=0) if e["dd_actual"] else "0", f"desde el pico de {usd(e['pico_equity'], dec=0)}", "neg" if e["dd_actual"] else "neu"),
        ("Factor de recuperación", "—" if e["recovery"] is None else f"{e['recovery']:.2f}", "neto ÷ drawdown máximo", "neu"),
        ("Racha ganadora máx.", f"{e['racha_gan']['n']} trades", usd(e["racha_gan"]["usd"], dec=0) + " USD" if e["racha_gan"]["n"] else "", "pos" if e["racha_gan"]["n"] else "neu"),
        ("Racha perdedora máx.", f"{e['racha_perd']['n']} trades", usd(e["racha_perd"]["usd"], dec=0) + " USD" if e["racha_perd"]["n"] else "", "neg" if e["racha_perd"]["n"] else "neu"),
        ("Mejor / peor trade", f"{e['mejor_trade']['usd']:+,.0f} / {e['peor_trade']['usd']:+,.0f}",
         f"{_f(e['mejor_trade']['fecha'])} · {_f(e['peor_trade']['fecha'])}", "neu"),
        ("Costos (comisiones)", "—" if e["comisiones"] is None else num(e["comisiones"], 0),
         "" if e["pct_comision_bruto"] is None else f"{e['pct_comision_bruto']:.1f}% del bruto", "neu"),
        ("Duración promedio", M.fmt_duracion(e["dur_prom_min"]),
         "" if not e["n_con_duracion"] else f"en {e['n_con_duracion']} trades importados", "neu"),
    ]


def columnas_ops(D):
    """[(clave, encabezado)] de la tabla de operaciones. Se omiten las vacías."""
    cols = [("fecha", "Fecha"), ("hora", "Hora"), ("activo", "Activo"), ("lado", "Lado"),
            ("contratos", "Cts"), ("setup", "Setup"), ("bruto", "Bruto"), ("costo", "Costo"), ("neto", "Neto"),
            ("empresa", "Empresa"), ("cuenta", "Cuenta"), ("emo", "Emo")]
    ops = D.get("ops", [])
    return [(k, h) for k, h in cols if any(o.get(k) not in (None, "") for o in ops)]


def nombre_sugerido(D, ext):
    ops = D.get("ops") or []
    if ops:
        rango = f"{ops[0]['fecha']}_a_{ops[-1]['fecha']}"
    else:
        rango = datetime.now().strftime("%Y-%m-%d")
    return f"registro-operativa_{rango}{ext}"


# ---------------------------------------------------------------- escritura segura
def _escribir(ruta, escritor, modo="wb", **kw):
    """Escribe en un temporal y recién al final reemplaza el destino, así un error
    nunca deja un archivo a medias ni pisa el que el usuario ya tenía."""
    carpeta = os.path.dirname(os.path.abspath(ruta)) or "."
    try:
        fd, tmp = tempfile.mkstemp(dir=carpeta, suffix=".tmp")
    except OSError as e:
        raise ExportError(f"No pude escribir en esa carpeta: {e.strerror or e}")
    os.close(fd)
    try:
        escritor(tmp)
        os.replace(tmp, ruta)
    except PermissionError:
        raise ExportError("No pude guardar el archivo. Si lo tenés abierto en otro programa "
                          "(por ejemplo Excel o un visor de PDF), cerralo y probá de nuevo.")
    except OSError as e:
        raise ExportError(f"No pude guardar el archivo: {e.strerror or e}")
    finally:
        if os.path.exists(tmp):
            try:
                os.remove(tmp)
            except OSError:
                pass
    return ruta


# ---------------------------------------------------------------- CSV
def a_csv(D, ruta):
    sep = D["opciones"].get("csv_sep", ",")
    cols = columnas_ops(D) or [("fecha", "Fecha")]
    if D["opciones"].get("comentarios") and any(o.get("comentario") for o in D.get("ops", [])):
        cols = cols + [("comentario", "Comentario")]

    def escribir(tmp):
        with open(tmp, "w", newline="", encoding="utf-8-sig") as fh:
            w = csv.writer(fh, delimiter=sep)
            w.writerow([h for _, h in cols])
            for o in D.get("ops", []):
                w.writerow(["" if o.get(k) is None else o.get(k) for k, _ in cols])
    return _escribir(ruta, escribir)


# ---------------------------------------------------------------- Excel
def _fecha_date(s):
    try:
        return datetime.strptime(s, "%Y-%m-%d").date()
    except (ValueError, TypeError):
        return s


def _hora_time(s):
    try:
        return datetime.strptime((s or "")[:5], "%H:%M").time()
    except ValueError:
        return s or None


def a_xlsx(D, ruta):
    try:
        from openpyxl import Workbook
        from openpyxl.chart import LineChart, Reference
        from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError:
        raise ExportError("Falta el componente de Excel (openpyxl). Reinstalá la aplicación.")
    sec = D["secciones"]
    if not sec:
        raise ExportError("Elegí al menos una sección para exportar.")
    if not D["n"]:
        raise ExportError("No hay operaciones en ese período.")

    INK, PAPER, SIGNAL, LINE = "0E1116", "F5F1EA", "2C5AA0", "DDD6CB"
    f_head = Font(name="Calibri", bold=True, color=PAPER)
    fill_head = PatternFill("solid", fgColor=INK)
    f_title = Font(name="Georgia", size=18, color=INK)
    f_muted = Font(name="Calibri", size=10, color="5C6671")
    f_bold = Font(name="Calibri", bold=True, color=INK)
    borde = Border(bottom=Side(style="thin", color=LINE))
    FMT_USD, FMT_USD0 = '#,##0.00;[Red]-#,##0.00', '#,##0;[Red]-#,##0'

    wb = Workbook()
    wb.properties.title = D["titulo"]
    wb.properties.creator = "Registro de operativa"
    wb.properties.subject = D["periodo"]
    wb.remove(wb.active)

    def cabecera(ws, fila, titulos, anchos=None):
        for j, t in enumerate(titulos, 1):
            c = ws.cell(row=fila, column=j, value=t)
            c.font, c.fill = f_head, fill_head
            c.alignment = Alignment(horizontal="center" if j > 1 else "left", vertical="center", wrap_text=True)
        if anchos:
            for j, w in enumerate(anchos, 1):
                ws.column_dimensions[get_column_letter(j)].width = w

    def encabezado_hoja(ws, subtitulo):
        ws["A1"] = D["titulo"]
        ws["A1"].font = f_title
        ws["A2"] = " · ".join(x for x in (D["periodo"], D["filtros"], subtitulo) if x)
        ws["A2"].font = f_muted
        ws["A3"] = f"Generado el {D['generado']} · Registro de operativa v{D['version']}"
        ws["A3"].font = f_muted

    # ---- Resumen
    if "resumen" in sec or "riesgo" in sec:
        ws = wb.create_sheet("Resumen")
        encabezado_hoja(ws, "")
        fila = 5
        bloques = []
        if "resumen" in sec:
            bloques.append(("Resumen", kpis_resumen(D)))
        if "riesgo" in sec:
            bloques.append(("Riesgo, rachas y costos", kpis_riesgo(D)))
        for titulo, items in bloques:
            ws.cell(row=fila, column=1, value=titulo).font = Font(name="Georgia", size=13, color=INK)
            fila += 1
            cabecera(ws, fila, ["Métrica", "Valor", "Detalle"])
            for lab, val, det, t in items:
                fila += 1
                ws.cell(row=fila, column=1, value=lab).font = f_bold
                c = ws.cell(row=fila, column=2, value=val)
                c.alignment = Alignment(horizontal="right")
                c.font = Font(name="Calibri", bold=True,
                              color={"pos": "3A6B3A", "neg": "8B3A2E"}.get(t, INK))
                ws.cell(row=fila, column=3, value=det).font = f_muted
                for j in (1, 2, 3):
                    ws.cell(row=fila, column=j).border = borde
            fila += 2
        for j, w in enumerate((30, 22, 46), 1):
            ws.column_dimensions[get_column_letter(j)].width = w
        ws.sheet_view.showGridLines = False

    # ---- Operaciones
    if "ops" in sec:
        ws = wb.create_sheet("Operaciones")
        cols = columnas_ops(D)
        cabecera(ws, 1, [h for _, h in cols])
        for i, o in enumerate(D["ops"], 2):
            for j, (k, _) in enumerate(cols, 1):
                v = o.get(k)
                if k == "fecha":
                    v = _fecha_date(v)
                elif k == "hora":
                    v = _hora_time(v)
                c = ws.cell(row=i, column=j, value=v)
                if k == "fecha":
                    c.number_format = "dd/mm/yyyy"
                elif k == "hora":
                    c.number_format = "hh:mm"
                elif k in ("bruto", "costo", "neto"):
                    c.number_format = FMT_USD
                elif k == "contratos":
                    c.number_format = "0.##"
        n = len(D["ops"])
        ref = f"A1:{get_column_letter(len(cols))}{n + 1}"
        t = Table(displayName="Operaciones", ref=ref)
        t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
        ws.add_table(t)
        ws.freeze_panes = "A2"
        anchos = dict(fecha=12, hora=8, activo=11, lado=8, contratos=7, setup=18, bruto=12, costo=10,
                      neto=12, empresa=18, cuenta=16, emo=6)
        for j, (k, _) in enumerate(cols, 1):
            ws.column_dimensions[get_column_letter(j)].width = anchos.get(k, 14)
        # totales que respetan el filtro de la tabla (SUBTOTAL ignora lo oculto)
        fila_t = n + 3
        ws.cell(row=fila_t, column=1, value="Totales (según filtro)").font = f_bold
        for j, (k, _) in enumerate(cols, 1):
            L = get_column_letter(j)
            if k in ("bruto", "costo", "neto"):
                c = ws.cell(row=fila_t, column=j, value=f"=SUBTOTAL(109,{L}2:{L}{n + 1})")
                c.number_format, c.font = FMT_USD, f_bold
            elif k == "activo":
                c = ws.cell(row=fila_t, column=j, value=f"=SUBTOTAL(103,{L}2:{L}{n + 1})")
                c.font = f_bold
                ws.cell(row=fila_t + 1, column=j, value="trades").font = f_muted

    # ---- Por día
    if "dias" in sec:
        ws = wb.create_sheet("Por día")
        cabecera(ws, 1, ["Fecha", "Día", "Trades", "Neto", "Win rate", "Acumulado"], (12, 7, 8, 12, 10, 13))
        for i, d in enumerate(reversed(D["dias"]), 2):
            vals = [_fecha_date(d["label"]), d["dia"], d["n"], d["neto"], d["wr"] / 100.0, d["acum"]]
            for j, v in enumerate(vals, 1):
                c = ws.cell(row=i, column=j, value=v)
                c.number_format = {1: "dd/mm/yyyy", 4: FMT_USD, 5: "0%", 6: FMT_USD}.get(j, "General")
        ws.freeze_panes = "A2"

    # ---- Análisis (hora / activo / semana / setup)
    grupos = [(k, t, D.get(k)) for k, t in (("hora", "Por hora de entrada"), ("activo", "Por activo"),
              ("semana", "Por día de la semana"), ("setup", "Por setup (playbook)")) if k in sec and D.get(k)]
    if grupos:
        ws = wb.create_sheet("Análisis")
        fila = 1
        for _, titulo, items in grupos:
            ws.cell(row=fila, column=1, value=titulo).font = Font(name="Georgia", size=13, color=INK)
            fila += 1
            cabecera(ws, fila, ["Grupo", "Trades", "Neto", "Win rate", "Promedio", "Mejor", "Peor"],
                     (22, 9, 13, 10, 12, 12, 12))
            for g in items:
                fila += 1
                vals = [g["label"], g["n"], g["neto"], g["wr"] / 100.0, g["prom"], g["mejor"], g["peor"]]
                for j, v in enumerate(vals, 1):
                    c = ws.cell(row=fila, column=j, value=v)
                    if j in (3, 5, 6, 7):
                        c.number_format = FMT_USD
                    elif j == 4:
                        c.number_format = "0%"
            fila += 3
        ws.sheet_view.showGridLines = False

    # ---- Curva (datos + gráfico nativo)
    if "curva" in sec:
        ws = wb.create_sheet("Curva")
        cabecera(ws, 1, ["N.º de trade", "Fecha", "Equity (USD)"], (12, 12, 14))
        for i, (f, v) in enumerate(zip(D["curva_fechas"], D["curva"]), 2):
            ws.cell(row=i, column=1, value=i - 1)
            c = ws.cell(row=i, column=2, value=_fecha_date(f))
            c.number_format = "dd/mm/yyyy"
            ws.cell(row=i, column=3, value=v).number_format = FMT_USD
        ch = LineChart()
        ch.title = "Curva de equity"
        ch.y_axis.title = "USD acumulados"
        ch.x_axis.title = "N.º de trade"
        ch.x_axis.delete = False
        ch.y_axis.delete = False
        ch.legend = None
        ch.height, ch.width = 9, 22
        datos = Reference(ws, min_col=3, min_row=1, max_row=len(D["curva"]) + 1)
        ch.add_data(datos, titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=len(D["curva"]) + 1))
        s = ch.series[0]
        s.graphicalProperties.line.solidFill = SIGNAL
        s.graphicalProperties.line.width = 22000
        s.smooth = False
        ws.add_chart(ch, "E2")
        ws.freeze_panes = "A2"

    if D["opciones"].get("marca") and wb.worksheets:
        ws0 = wb.worksheets[0]
        ws0["A4"] = MARCA_TEXTO
        ws0["A4"].font = Font(name="Calibri", size=9, italic=True, color=SIGNAL)

    return _escribir(ruta, lambda tmp: wb.save(tmp))


# ---------------------------------------------------------------- HTML
def _svg_equity(vals, fechas, w=760, h=260):
    """Curva de equity como SVG con título y descripción (lectores de pantalla)."""
    import html as H
    if len(vals) < 2:
        return ""
    pl, pr, pt, pb = 62, 14, 14, 26
    pts = vals if len(vals) <= 800 else [vals[int(i * (len(vals) - 1) / 799)] for i in range(800)]
    lo, hi = min(0.0, min(pts)), max(0.0, max(pts))
    if hi == lo:
        hi, lo = hi + 1, lo - 1
    X = lambda i: pl + (w - pl - pr) * i / (len(pts) - 1)
    Y = lambda v: pt + (h - pt - pb) * (1 - (v - lo) / (hi - lo))
    path = " ".join(f"{X(i):.1f},{Y(v):.1f}" for i, v in enumerate(pts))
    area = f"{X(0):.1f},{Y(0):.1f} {path} {X(len(pts) - 1):.1f},{Y(0):.1f}"
    desc = (f"Equity acumulada en {len(vals)} trades, de {_f(fechas[0])} a {_f(fechas[-1])}. "
            f"Cierra en {usd(vals[-1], dec=0)} USD, con un máximo de {usd(max(vals), dec=0)} y un mínimo de {usd(min(vals), dec=0)}.")
    lab = lambda v: f"{v:+,.0f}" if v else "0"
    return (
        f'<svg class="equity" viewBox="0 0 {w} {h}" role="img" aria-labelledby="eq-t eq-d" preserveAspectRatio="xMidYMid meet">'
        f'<title id="eq-t">Curva de equity</title><desc id="eq-d">{H.escape(desc)}</desc>'
        f'<line x1="{pl}" x2="{w - pr}" y1="{Y(0):.1f}" y2="{Y(0):.1f}" class="cero"/>'
        f'<polygon points="{area}" class="area"/>'
        f'<polyline points="{path}" class="linea" fill="none"/>'
        f'<circle cx="{X(len(pts) - 1):.1f}" cy="{Y(pts[-1]):.1f}" r="4" class="punto"/>'
        f'<text x="{pl - 8}" y="{Y(hi) + 4:.1f}" text-anchor="end">{lab(hi)}</text>'
        f'<text x="{pl - 8}" y="{Y(0) + 4:.1f}" text-anchor="end">0</text>'
        + (f'<text x="{pl - 8}" y="{Y(lo) + 4:.1f}" text-anchor="end">{lab(lo)}</text>' if lo < 0 else "")
        + f'<text x="{pl}" y="{h - 6}" text-anchor="start">{_f(fechas[0])}</text>'
        f'<text x="{w - pr}" y="{h - 6}" text-anchor="end">{_f(fechas[-1])}</text></svg>')


def _fuentes_css():
    import base64
    from .paths import resource_path
    out = []
    for fam, arch in (("DM Sans", "DMSans-Regular.ttf"), ("Instrument Serif", "InstrumentSerif-Regular.ttf"),
                      ("JetBrains Mono", "JetBrainsMono-Regular.ttf")):
        try:
            with open(resource_path("assets", "fonts", arch), "rb") as fh:
                b64 = base64.b64encode(fh.read()).decode("ascii")
            out.append(f"@font-face{{font-family:'{fam}';src:url(data:font/ttf;base64,{b64}) format('truetype');font-display:swap}}")
        except OSError:
            pass
    return "\n".join(out)


_LOGO = ('<svg class="logo" viewBox="0 0 32 20" width="30" height="19" aria-hidden="true" focusable="false">'
         '<line x1="3" y1="16" x2="29" y2="16" stroke="currentColor" stroke-width="2" stroke-linecap="round"/>'
         '<line x1="21" y1="16" x2="21" y2="5" stroke="var(--signal)" stroke-width="2" stroke-linecap="round"/>'
         '<circle cx="21" cy="4" r="3" fill="var(--signal)"/></svg>')

_CSS = """
:root{--bg:%(bg)s;--card:%(card)s;--text:%(text)s;--muted:%(muted)s;--line:%(line)s;--signal:%(signal)s;--pos:%(pos)s;--neg:%(neg)s;
--serif:'Instrument Serif',Georgia,serif;--sans:'DM Sans','Segoe UI',system-ui,sans-serif;--mono:'JetBrains Mono',Consolas,ui-monospace,monospace}
*{box-sizing:border-box}html{-webkit-text-size-adjust:100%%}
body{margin:0;background:var(--bg);color:var(--text);font:15px/1.6 var(--sans)}
.skip{position:absolute;left:-999px;top:8px;background:var(--text);color:var(--bg);padding:8px 12px;border-radius:4px;z-index:9}
.skip:focus{left:8px}
a{color:var(--signal)}a:focus-visible,button:focus-visible,input:focus-visible{outline:3px solid var(--signal);outline-offset:2px}
.wrap{max-width:1000px;margin:0 auto;padding:40px 24px 64px}
header.top{display:flex;gap:14px;align-items:center;margin-bottom:6px;color:var(--text)}
.kicker{font:11px var(--mono);letter-spacing:.2em;text-transform:uppercase;color:var(--muted)}
h1{font:400 44px/1.1 var(--serif);margin:.15em 0 .1em}
h2{font:400 28px/1.2 var(--serif);margin:0 0 14px;padding-top:6px}
section{margin:44px 0 0}
.meta{color:var(--muted);font-size:13px;margin:0}
nav.toc{display:flex;flex-wrap:wrap;gap:6px 18px;margin:22px 0 0;padding:12px 0;border-top:1px solid var(--line);border-bottom:1px solid var(--line)}
nav.toc a{font:11px var(--mono);letter-spacing:.14em;text-transform:uppercase;color:var(--muted);text-decoration:none}
nav.toc a:hover{color:var(--text);text-decoration:underline}
dl.kpis{display:grid;grid-template-columns:repeat(auto-fit,minmax(200px,1fr));gap:12px;margin:0}
dl.kpis>div{background:var(--card);border:1px solid var(--line);border-radius:10px;padding:14px 16px}
dt{font:10px var(--mono);letter-spacing:.16em;text-transform:uppercase;color:var(--muted)}
dd{margin:0}dd.val{font:400 28px/1.2 var(--serif);margin-top:4px}dd.sub{font:12px var(--mono);color:var(--muted);margin-top:2px;min-height:1.2em}
.pos{color:var(--pos)}.neg{color:var(--neg)}
figure{margin:0;background:var(--card);border:1px solid var(--line);border-radius:10px;padding:12px}
svg.equity{width:100%%;height:auto;display:block}svg.equity text{font:11px var(--mono);fill:var(--muted)}
svg.equity .linea{stroke:var(--signal);stroke-width:2;stroke-linejoin:round}svg.equity .area{fill:var(--signal);opacity:.08}
svg.equity .cero{stroke:var(--muted);stroke-dasharray:4 4;opacity:.6}svg.equity .punto{fill:var(--signal)}
.tabla{overflow-x:auto;border:1px solid var(--line);border-radius:10px;background:var(--card)}
table{border-collapse:collapse;width:100%%;font-size:13.5px}
caption{text-align:left;padding:10px 14px;color:var(--muted);font-size:12px}
th,td{padding:8px 12px;text-align:right;border-top:1px solid var(--line);white-space:nowrap}
th:first-child,td:first-child,td.l,th.l{text-align:left}
thead th{background:var(--text);color:var(--bg);font:11px var(--mono);letter-spacing:.1em;text-transform:uppercase;border-top:0;position:sticky;top:0}
thead th button{all:unset;cursor:pointer;display:block;width:100%%;text-align:inherit;padding:0}
thead th button:focus-visible{outline:3px solid var(--signal);outline-offset:2px}
th[aria-sort=ascending] button::after{content:" ▲"}th[aria-sort=descending] button::after{content:" ▼"}
td.num{font-family:var(--mono)}td.com{white-space:normal;color:var(--muted);font-style:italic;text-align:left}
td .bar{display:block;height:4px;border-radius:2px;margin-top:4px;background:var(--signal);opacity:.55}
td .bar.neg{background:var(--neg)}td .bar.pos{background:var(--pos)}
.filtro{margin:0 0 10px}.filtro label{font:11px var(--mono);letter-spacing:.14em;text-transform:uppercase;color:var(--muted);margin-right:8px}
.filtro input{font:14px var(--sans);padding:7px 10px;border:1px solid var(--line);border-radius:6px;background:var(--card);color:var(--text);min-width:240px}
footer{margin-top:56px;padding-top:16px;border-top:1px solid var(--line);font:11px var(--mono);letter-spacing:.06em;color:var(--muted);display:flex;justify-content:space-between;flex-wrap:wrap;gap:8px}
@media (max-width:640px){h1{font-size:34px}.wrap{padding:24px 16px 48px}}
@media print{:root{--bg:#fff;--card:#fff;--text:#0E1116;--muted:#5C6671;--line:#DDD6CB;--signal:#2C5AA0;--pos:#3A6B3A;--neg:#8B3A2E}
body{font-size:12px}.wrap{padding:0}nav.toc,.filtro,.skip{display:none}section{break-inside:avoid-page;margin-top:28px}tr{break-inside:avoid}
thead{display:table-header-group}thead th{position:static;-webkit-print-color-adjust:exact;print-color-adjust:exact}
.tabla{overflow:visible;border-radius:0}}
"""

_JS = """
(function(){
var t=document.getElementById('tabla-ops');if(!t)return;
var tb=t.tBodies[0],filas=[].slice.call(tb.rows).filter(function(r){return !r.classList.contains('com-row')});
function val(r,i){var c=r.cells[i];return c.getAttribute('data-v')||c.textContent.trim()}
[].forEach.call(t.tHead.rows[0].cells,function(th,i){
 var b=th.querySelector('button');if(!b)return;
 b.addEventListener('click',function(){
  var asc=th.getAttribute('aria-sort')!=='ascending';
  [].forEach.call(t.tHead.rows[0].cells,function(x){x.removeAttribute('aria-sort')});
  th.setAttribute('aria-sort',asc?'ascending':'descending');
  filas.sort(function(a,c){var x=val(a,i),y=val(c,i),nx=parseFloat(x),ny=parseFloat(y);
   var r=(!isNaN(nx)&&!isNaN(ny))?nx-ny:x.localeCompare(y,'es');return asc?r:-r});
  filas.forEach(function(r){tb.appendChild(r);if(r._com)tb.appendChild(r._com)});
 });
});
filas.forEach(function(r){var n=r.nextElementSibling;if(n&&n.classList.contains('com-row'))r._com=n});
var q=document.getElementById('filtro-ops'),st=document.getElementById('estado-filtro');
if(q)q.addEventListener('input',function(){var s=q.value.toLowerCase(),n=0;
 filas.forEach(function(r){var ok=r.textContent.toLowerCase().indexOf(s)>-1;r.hidden=!ok;if(r._com)r._com.hidden=!ok;if(ok)n++});
 st.textContent=n+' de '+filas.length+' operaciones'});
})();
"""


def a_html(D, ruta):
    import html as H
    e = H.escape
    if not D["secciones"]:
        raise ExportError("Elegí al menos una sección para exportar.")
    if not D["n"]:
        raise ExportError("No hay operaciones en ese período.")
    sec, op, partes, toc = D["secciones"], D["opciones"], [], []

    def seccion(clave, titulo, cuerpo):
        toc.append(f'<a href="#s-{clave}">{e(titulo)}</a>')
        partes.append(f'<section id="s-{clave}" aria-labelledby="h-{clave}"><h2 id="h-{clave}">{e(titulo)}</h2>{cuerpo}</section>')

    def kpis(items):
        return '<dl class="kpis">' + "".join(
            f'<div><dt>{e(l)}</dt><dd class="val {t}">{e(v)}</dd><dd class="sub">{e(d)}</dd></div>'
            for l, v, d, t in items) + "</dl>"

    def tabla_grupo(items, titulo):
        mx = max((abs(g["neto"]) for g in items), default=1) or 1
        filas = "".join(
            f'<tr><td class="l">{e(str(g["label"]))}</td><td class="num">{g["n"]}</td>'
            f'<td class="num {tono(g["neto"])}">{usd(g["neto"], dec=0)}'
            f'<span class="bar {tono(g["neto"])}" style="width:{abs(g["neto"]) / mx * 100:.0f}%" aria-hidden="true"></span></td>'
            f'<td class="num">{g["wr"]:.0f}%</td><td class="num">{usd(g["prom"], dec=1)}</td></tr>' for g in items)
        return (f'<div class="tabla"><table><caption>{e(titulo)}. Las barras comparan el neto entre filas; '
                f'con pocos trades por fila, no saques conclusiones.</caption><thead><tr><th class="l" scope="col">Grupo</th>'
                f'<th scope="col">Trades</th><th scope="col">Neto</th><th scope="col">Win rate</th><th scope="col">Promedio</th></tr></thead>'
                f'<tbody>{filas}</tbody></table></div>')

    if "resumen" in sec:
        seccion("resumen", "Resumen y métricas clave", kpis(kpis_resumen(D)))
    if "riesgo" in sec:
        seccion("riesgo", "Riesgo: drawdown, rachas y costos", kpis(kpis_riesgo(D)))
    if "curva" in sec:
        svg = _svg_equity(D["curva"], D["curva_fechas"])
        seccion("curva", "Curva de equity",
                f'<figure>{svg}</figure>' if svg else '<p class="meta">Hacen falta al menos 2 trades para dibujar la curva.</p>')
    if "dias" in sec:
        mx = max((abs(d["neto"]) for d in D["dias"]), default=1) or 1
        filas = "".join(
            f'<tr><td class="l">{_f(d["label"])} <span class="meta">{e(d["dia"])}</span></td><td class="num">{d["n"]}</td>'
            f'<td class="num {tono(d["neto"])}">{usd(d["neto"], dec=0)}<span class="bar {tono(d["neto"])}" '
            f'style="width:{abs(d["neto"]) / mx * 100:.0f}%" aria-hidden="true"></span></td>'
            f'<td class="num">{d["wr"]:.0f}%</td><td class="num {tono(d["acum"])}">{usd(d["acum"], dec=0)}</td></tr>' for d in D["dias"])
        seccion("dias", "Resultado por día",
                '<div class="tabla"><table><caption>Del día más nuevo al más viejo.</caption><thead><tr><th class="l" scope="col">Día</th>'
                '<th scope="col">Trades</th><th scope="col">Neto</th><th scope="col">Win rate</th><th scope="col">Acumulado</th></tr></thead>'
                f'<tbody>{filas}</tbody></table></div>')
    for clave, titulo in (("hora", "Por hora de entrada"), ("activo", "Por activo"),
                          ("semana", "Por día de la semana"), ("setup", "Por setup (playbook)")):
        if clave in sec and D.get(clave):
            seccion(clave, titulo, tabla_grupo(D[clave], titulo))
    if "ops" in sec:
        cols = columnas_ops(D)
        ths = "".join(
            f'<th scope="col" class="{"l" if k in ("fecha", "hora", "activo", "lado", "setup", "empresa", "cuenta") else ""}">'
            f'<button type="button">{e(h)}</button></th>' for k, h in cols)
        filas = []
        for o in D["ops"]:
            celdas = []
            for k, _ in cols:
                v = o.get(k)
                if k == "fecha":
                    celdas.append(f'<td class="l" data-v="{e(v)}">{_f(v)}</td>')
                elif k in ("bruto", "costo", "neto"):
                    celdas.append(f'<td class="num {tono(v) if k == "neto" else ""}" data-v="{"" if v is None else v}">{usd(v)}</td>')
                elif k in ("contratos", "emo"):
                    celdas.append(f'<td class="num">{"" if v is None else f"{v:g}"}</td>')
                else:
                    celdas.append(f'<td class="l">{e(str(v or ""))}</td>')
            filas.append("<tr>" + "".join(celdas) + "</tr>")
            if o.get("comentario"):
                filas.append(f'<tr class="com-row"><td class="com" colspan="{len(cols)}">{e(o["comentario"])}</td></tr>')
        seccion("ops", "Detalle de operaciones",
                '<div class="filtro"><label for="filtro-ops">Filtrar</label><input id="filtro-ops" type="search" '
                'placeholder="activo, setup, fecha…" autocomplete="off"> <span id="estado-filtro" class="meta" role="status" aria-live="polite"></span></div>'
                f'<div class="tabla"><table id="tabla-ops"><caption>{len(D["ops"])} operaciones. Hacé click en un encabezado para ordenar.</caption>'
                f'<thead><tr>{ths}</tr></thead><tbody>{"".join(filas)}</tbody></table></div>')

    meta = " · ".join(x for x in (D["periodo"], D["filtros"], f"{D['n']} trades") if x)
    marca = f'<span>{e(MARCA_TEXTO)}</span>' if op.get("marca") else "<span></span>"
    css = (_fuentes_css() if op.get("fuentes", True) else "") + _CSS % D["paleta"]
    doc = (f'<!doctype html><html lang="es"><head><meta charset="utf-8">'
           f'<meta name="viewport" content="width=device-width,initial-scale=1">'
           f'<meta name="generator" content="Registro de operativa {e(D["version"])}"><title>{e(D["titulo"])}</title>'
           f'<style>{css}</style></head><body><a class="skip" href="#contenido">Saltar al contenido</a>'
           f'<div class="wrap"><header class="top">{_LOGO}<span class="kicker">Registro de operativa</span></header>'
           f'<h1>{e(D["titulo"])}</h1><p class="meta">{e(meta)}</p>'
           f'<nav class="toc" aria-label="Secciones">{"".join(toc)}</nav><main id="contenido">{"".join(partes)}</main>'
           f'<footer><span>Generado el {e(D["generado"])}</span>{marca}</footer></div><script>{_JS}</script></body></html>')

    def escribir(tmp):
        with open(tmp, "w", encoding="utf-8") as fh:
            fh.write(doc)
    return _escribir(ruta, escribir)


# ---------------------------------------------------------------- despachador
def exportar(formato, D, ruta):
    """Escribe el reporte en `ruta`. formato: 'pdf' | 'html' | 'xlsx' | 'csv'."""
    if formato == "pdf":
        from .exportar_pdf import a_pdf
        return a_pdf(D, ruta)
    if formato == "html":
        return a_html(D, ruta)
    if formato == "xlsx":
        return a_xlsx(D, ruta)
    if formato == "csv":
        if not D["n"]:
            raise ExportError("No hay operaciones en ese período.")
        return a_csv(D, ruta)
    raise ExportError(f"Formato desconocido: {formato}")
